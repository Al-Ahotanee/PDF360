import io

import pytest

from app.services.pdf_engine import conversion as conv
from app.services.pdf_engine.core import PDFEngineError, get_page_count


@pytest.fixture
def sample_docx() -> bytes:
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_heading("Test Report", level=1)
    doc.add_paragraph("Body paragraph for conversion testing.")
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def pdf_with_gridded_table() -> bytes:
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf)
    data = [["Name", "Score"], ["Alice", "90"], ["Bob", "85"]]
    table = Table(data)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
    doc.build([table])
    return buf.getvalue()


def test_office_to_pdf(sample_docx):
    pdf_bytes = conv.office_to_pdf(sample_docx, ".docx")
    assert pdf_bytes[:4] == b"%PDF"


def test_images_to_pdf(sample_image):
    pdf_bytes = conv.images_to_pdf([sample_image, sample_image])
    assert get_page_count(pdf_bytes) == 2


def test_images_to_pdf_requires_at_least_one():
    with pytest.raises(PDFEngineError):
        conv.images_to_pdf([])


def test_pdf_to_images(sample_docx):
    pdf_bytes = conv.office_to_pdf(sample_docx, ".docx")
    images = conv.pdf_to_images(pdf_bytes, fmt="png")
    assert len(images) == get_page_count(pdf_bytes)
    assert images[0][:8] == b"\x89PNG\r\n\x1a\n"


def test_markdown_to_pdf():
    pdf_bytes = conv.markdown_to_pdf("# Title\n\nSome **bold** text.\n\n- one\n- two")
    assert pdf_bytes[:4] == b"%PDF"


def test_html_to_pdf():
    pdf_bytes = conv.html_to_pdf("<html><body><p>Hello world.</p></body></html>")
    assert pdf_bytes[:4] == b"%PDF"
    assert "Hello world" in conv.pdf_to_txt(pdf_bytes)


def test_pdf_to_docx_reconstruction(sample_docx):
    pdf_bytes = conv.office_to_pdf(sample_docx, ".docx")
    docx_out = conv.pdf_to_docx(pdf_bytes)

    from docx import Document as ReadDocx

    reconstructed = ReadDocx(io.BytesIO(docx_out))
    full_text = "\n".join(p.text for p in reconstructed.paragraphs)
    assert "Test Report" in full_text


def test_pdf_to_pptx_slide_per_page(sample_docx):
    pdf_bytes = conv.office_to_pdf(sample_docx, ".docx")
    pptx_out = conv.pdf_to_pptx(pdf_bytes)

    from pptx import Presentation

    prs = Presentation(io.BytesIO(pptx_out))
    assert len(list(prs.slides)) == get_page_count(pdf_bytes)


def test_pdf_to_xlsx_extracts_gridded_table(pdf_with_gridded_table):
    xlsx_out = conv.pdf_to_xlsx(pdf_with_gridded_table)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_out))
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    assert rows[0] == ("Name", "Score")
    assert rows[1] == ("Alice", "90")


def test_pdf_to_xlsx_raises_with_no_tables(sample_docx):
    pdf_bytes = conv.office_to_pdf(sample_docx, ".docx")
    with pytest.raises(PDFEngineError):
        conv.pdf_to_xlsx(pdf_bytes)
